"""ゲームデータ名の韓国語辞書(`apps/desktop/src/i18n/ko/names.json`)を作る(docs/adr/022-i18n.md 段階 4)。

出典が確実なものだけ入れる。出典・照合方法:

1. 装備名: 韓国コミュニティ装備整理シート(Google スプレッドシート
   `1rT24bRdfsqcX3N4JbRx1dZhqyEAf5OcAPnwPerd18Ds`、sacred_kr.rs と同じ出典)の武器・防具タブを、
   9 値(values_min/max)の完全一致 + 系列名の音の対応で日本 Tale Wiki 名へ照合する。
   候補が 0 件・複数残る・系列が合わないものは入れずに報告する。
2. キャラ名: 韓国公式のキャラ一覧 `https://tales.nexon.com/About/Character` の 19 人と
   crates/gamedata/src/characters.rs の 19 人を id で突き合わせる。
3. 装備アビリティ名(シートの「어빌리티」タブ): 確実な照合キーが無いため今回は見送り(下記参照)。

前提: `cargo test -p gamedata --test dump_names -- --ignored dump_names` で
`tools/i18n/out/*.json`(統合後の日本語名。gitignore)を作ってあること。このスクリプトは
無ければ自動で実行する。

使い方: `python tools/i18n/import_names.py`
"""

from __future__ import annotations

import json
import re
import subprocess
import sys
import urllib.request
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
OUT_DIR = ROOT / "tools" / "i18n" / "out"
CACHE_DIR = ROOT / "tools" / "i18n" / "cache"
NAMES_JSON = ROOT / "apps" / "desktop" / "src" / "i18n" / "ko" / "names.json"

SHEET_ID = "1rT24bRdfsqcX3N4JbRx1dZhqyEAf5OcAPnwPerd18Ds"
SHEET_URL = f"https://docs.google.com/spreadsheets/d/{SHEET_ID}/export?format=xlsx"
SHEET_CACHE = CACHE_DIR / "sheet.xlsx"
CHAR_URL = "https://tales.nexon.com/About/Character"
CHAR_CACHE = CACHE_DIR / "characters.html"
USER_AGENT = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
    "(KHTML, like Gecko) Chrome/120.0 Safari/537.36"
)

# EquipmentValues の並び(crates/domain/src/equipment.rs)。シートは
# …명중,회피,민첩함,크리 の順であり、accuracy の後ろが critical/evasion/agility の順とずれる
# (docs/damage-formula.md、旧 sacred_kr.rs でも同じ並べ替えを行っている)。
STAT_ORDER = [
    "thrust", "slash", "physical_defense", "magic_attack", "magic_defense",
    "accuracy", "critical", "evasion", "agility",
]
# シートの列見出し → EquipmentValues のフィールド名
HEADER_LABELS = {
    "찌르기": "thrust",
    "베기": "slash",
    "방어": "physical_defense",
    "마공": "magic_attack",
    "마방": "magic_defense",
    "명중": "accuracy",
    "크리": "critical",
    "회피": "evasion",
    "민첩함": "agility",
}

# 系列名の音の対応(韓国語の頭 → 日本語の頭)。長い方から試す(改-세크리드 が 세크리드 より先)。
# 実データで確かめた対応(2026-09-24)。
SERIES_MAP = {
    "改-세크리드": "改・セイクリッド",
    "세크리드": "セイクリッド",
    "아노마라드 공화국": "アノマラド共和国",
    "아노마라드 왕국": "アノマラド王国",
    "데모닉": "デモニック",
    "인퍼널": "インファーナル",
    "아퀼루스": "アクィルス",
    "어비스": "アビス",
    "이클립스": "エクリプス",
}
SERIES_ORDER = sorted(SERIES_MAP, key=len, reverse=True)

# 武器・防具の数値タブ。目次・アビリティ・拡張表・レリック(神鳥のブレスレット等、命名規則が
# 別)は対象外。펜듈럼・핸드런처・아밍소드・소드셰이프・해머・물리탄창・마법탄창 は配置が特殊で
# 汎用パーサでは 0 件(試作で確認済み)。読めた分だけ採る。
ITEM_TABS = [
    "세검", "장검", "평도", "대검", "태도", "스태프", "로드", "메이스", "단검", "단도", "도끼",
    "창", "봉", "채찍", "플레일", "스몰소드", "완드", "물리총", "마법총", "클로", "카라", "셉터",
    "핸드벨", "물리검", "마법검", "사이드", "해머", "토템", "핸드런처", "아밍소드", "소드셰이프",
    "펜듈럼", "물리탄창", "마법탄창", "물리검(sub)", "마법검(sub)",
    "로브", "아머", "메일", "마법갑옷", "슈츠",
    "아노마라드 공화국 시리즈", "아노마라드 왕국 시리즈",
    "데모닉 장비 세트", "인퍼널 장비 세트", "아퀼루스 장비 세트", "어비스 장비 세트",
    "이클립스 장비 세트", "세크리드 장비 세트", "改-세크리드 장비 세트",
    "방패", "리스트", "밴드", "암릿", "수정구", "스펠북",
]


def ensure_cache() -> None:
    CACHE_DIR.mkdir(parents=True, exist_ok=True)
    if not SHEET_CACHE.exists():
        print(f"シートを取得: {SHEET_URL}")
        urllib.request.urlretrieve(SHEET_URL, SHEET_CACHE)
    if not CHAR_CACHE.exists():
        print(f"キャラ一覧を取得: {CHAR_URL}")
        req = urllib.request.Request(CHAR_URL, headers={"User-Agent": USER_AGENT})
        with urllib.request.urlopen(req) as resp:
            CHAR_CACHE.write_bytes(resp.read())


def ensure_dump() -> None:
    if (OUT_DIR / "equipment_catalog.json").exists() and (OUT_DIR / "characters.json").exists():
        return
    print("crates/gamedata の名前を抽出: cargo test -p gamedata --test dump_names -- --ignored dump_names")
    subprocess.run(
        ["cargo", "test", "-p", "gamedata", "--test", "dump_names", "--", "--ignored", "dump_names"],
        cwd=ROOT, check=True,
    )


# ---- 装備名 ----

def parse_min_max(cell) -> tuple[int, int] | None:
    """シートの 1 セルを (min, max) に。'-'/空 は (0, 0)。日付型はシートの入力ミス(バグ)で
    数値として読めないので None(この項目は取り込まない)。"""
    if cell is None:
        return (0, 0)
    if isinstance(cell, bool):
        return None
    if isinstance(cell, (int, float)):
        return (int(cell), int(cell))
    if isinstance(cell, str):
        text = cell.strip()
        if text in ("", "-"):
            return (0, 0)
        if re.fullmatch(r"-?\d+", text):
            v = int(text)
            return (v, v)
        m = re.fullmatch(r"(-?\d+)\s*-\s*(-?\d+)", text)
        if m:
            return (int(m.group(1)), int(m.group(2)))
        return None
    return None  # datetime 等、シート側のバグ


def parse_tab(ws) -> list[dict]:
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    header_map: dict[str, int] = {}
    for idx, cell in enumerate(rows[0]):
        if isinstance(cell, str) and cell.strip() in HEADER_LABELS:
            header_map[HEADER_LABELS[cell.strip()]] = idx

    items: list[dict] = []
    pending_name: str | None = None
    for row in rows[1:]:
        first = row[0].strip() if isinstance(row[0], str) else None
        is_option_row = any(isinstance(c, str) and c.strip() == "기본 옵션" for c in row)
        if is_option_row:
            if pending_name is None:
                continue
            mins: list[int] = []
            maxs: list[int] = []
            ok = True
            for stat in STAT_ORDER:
                idx = header_map.get(stat)
                cell = row[idx] if idx is not None else None
                parsed = parse_min_max(cell)
                if parsed is None:
                    ok = False
                    break
                mins.append(parsed[0])
                maxs.append(parsed[1])
            if ok:
                items.append({"name": pending_name, "min": tuple(mins), "max": tuple(maxs)})
            pending_name = None
        elif first and not any(isinstance(c, str) and c.strip() == "한계치" for c in row):
            pending_name = first
    return items


def load_kr_items() -> list[tuple[str, list[dict]]]:
    import openpyxl

    wb = openpyxl.load_workbook(SHEET_CACHE, read_only=True, data_only=True)
    items: list[tuple[str, list[dict]]] = []
    for tab in ITEM_TABS:
        if tab not in wb.sheetnames:
            print(f"  (タブが無い: {tab})")
            continue
        found = parse_tab(wb[tab])
        for it in found:
            it["tab"] = tab
        items.append((tab, found))
    return items


def series_of(kr_name: str) -> tuple[str, str] | None:
    for kr_prefix in SERIES_ORDER:
        if kr_name.startswith(kr_prefix):
            return kr_prefix, SERIES_MAP[kr_prefix]
    return None


# 同じ 9 値を複数の武器種・部位が共有する組(双剣本体/Sub、カーラ/アーミングソード、
# ガントレット/ブーツ 等)がある。系列名だけでは絞れないので、系列の後ろの語で追加照合する
# (実データで確かめた分だけ。2026-09-24)。
KR_TYPE_TO_JP = {
    "카라": "カーラ",
    "아밍소드": "アーミングソード",
    "건틀렛": "ガントレット",
    "부츠": "ブーツ",
}


def narrow_candidates(tab: str, kr_name: str, series_prefix: str, candidates: list[str]) -> list[str]:
    """系列名だけで複数残るとき、タブ・語尾で絞れる分だけ絞る。絞れなければそのまま返す。"""
    if tab in ("물리검(sub)", "마법검(sub)"):
        narrowed = [c for c in candidates if "sub" in c.lower()]
        return narrowed or candidates
    if tab in ("물리검", "마법검"):
        narrowed = [c for c in candidates if "sub" not in c.lower()]
        return narrowed or candidates
    extra = kr_name[len(series_prefix):].strip()
    jp_word = KR_TYPE_TO_JP.get(extra)
    if jp_word:
        narrowed = [c for c in candidates if jp_word in c]
        return narrowed or candidates
    return candidates


def import_equipment(jp_catalog: list[dict], kr_by_tab: list[tuple[str, list[dict]]]) -> dict[str, str]:
    # (min, max) -> 候補(日本語名のリスト)
    jp_index: dict[tuple, list[str]] = {}
    for it in jp_catalog:
        key = (tuple(it["values_min"][s] for s in STAT_ORDER), tuple(it["values_max"][s] for s in STAT_ORDER))
        jp_index.setdefault(key, []).append(it["name"])

    result: dict[str, str] = {}
    counts = {"採用": 0, "候補なし": 0, "候補複数": 0, "系列不明": 0, "系列不一致": 0, "パース失敗": 0}
    conflicts: list[str] = []

    for tab, items in kr_by_tab:
        for it in items:
            key = (it["min"], it["max"])
            candidates = jp_index.get(key, [])
            if not candidates:
                counts["候補なし"] += 1
                continue
            series = series_of(it["name"])
            if series is None:
                counts["系列不明"] += 1
                continue
            kr_prefix, jp_prefix = series
            filtered = [c for c in candidates if c.lstrip("†").startswith(jp_prefix)]
            if len(filtered) > 1:
                filtered = narrow_candidates(tab, it["name"], kr_prefix, filtered)
            if len(filtered) == 0:
                counts["系列不一致"] += 1
                continue
            if len(filtered) > 1:
                counts["候補複数"] += 1
                continue
            jp_name = filtered[0]
            prefix = "†" if jp_name.startswith("†") else ""
            kr_value = prefix + it["name"]
            if jp_name in result and result[jp_name] != kr_value:
                conflicts.append(f"{jp_name!r}: {result[jp_name]!r} != {kr_value!r} ({tab})")
                continue
            result[jp_name] = kr_value
            counts["採用"] += 1

    if conflicts:
        print("衝突(同じ日本語名に違う訳):")
        for c in conflicts:
            print(f"  {c}")
        raise SystemExit(1)

    print(f"装備名: {counts}")
    return result


# ---- キャラ名 ----

# crates/gamedata/src/characters.rs の id → tales.nexon.com の URL スラッグ(2026-09-24 実データで確認)
CHARACTER_SLUGS = {
    "lucian": "Lucian", "boris": "Boris", "ispin": "Ispin", "maximin": "Maximin",
    "tichiel": "Tichiel", "nayatorei": "Naya", "siberin": "Sivelin", "mira": "Mila",
    "joshua": "Josua", "chloe": "Cloe", "ranjie": "Lanziee", "isaac": "Issac",
    "anais": "Anais", "isolet": "Isolet", "benya": "Benya", "roamini": "Roamini",
    "nocturne": "Nocturne", "leeche": "Clarice", "yefnen": "Yevgnen",
}


def import_characters(jp_characters: list[dict]) -> dict[str, str]:
    html = CHAR_CACHE.read_text(encoding="utf-8")
    by_slug: dict[str, str] = {}
    for m in re.finditer(
        r'href="/About/Character/([^"]+)".*?headline">([^<]+)</span>', html, re.S,
    ):
        slug, kr_full = m.groups()
        by_slug[slug] = kr_full.strip()

    result: dict[str, str] = {}
    missing: list[str] = []
    for char in jp_characters:
        slug = CHARACTER_SLUGS.get(char["id"])
        kr_full = by_slug.get(slug) if slug else None
        if not kr_full:
            missing.append(char["id"])
            continue
        # JP 側は姓を持たないので、韓国語も先頭の名だけを取る(例 "루시안 칼츠" -> "루시안")
        kr_first = kr_full.split()[0]
        result[char["name"]] = kr_first
    if missing:
        print(f"キャラ名: 照合できなかった id = {missing}")
    print(f"キャラ名: 採用 {len(result)} / 全 {len(jp_characters)}")
    return result


def main() -> int:
    sys.stdout.reconfigure(encoding="utf-8")  # type: ignore[attr-defined]
    ensure_cache()
    ensure_dump()

    jp_catalog = json.loads((OUT_DIR / "equipment_catalog.json").read_text(encoding="utf-8"))
    jp_characters = json.loads((OUT_DIR / "characters.json").read_text(encoding="utf-8"))

    print("装備整理シートを解析中…")
    kr_by_tab = load_kr_items()
    for tab, items in kr_by_tab:
        print(f"  {tab}: {len(items)} 件")

    names: dict[str, str] = {}
    names.update(import_equipment(jp_catalog, kr_by_tab))
    names.update(import_characters(jp_characters))

    print(
        "装備アビリティ名(어빌리티 タブ): "
        "確実な照合キー(id・client アイテム ID)が無いため見送り(docs/adr/022-i18n.md 参照)"
    )

    sorted_names = {k: names[k] for k in sorted(names)}
    NAMES_JSON.parent.mkdir(parents=True, exist_ok=True)
    NAMES_JSON.write_text(
        json.dumps(sorted_names, ensure_ascii=False, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )
    print(f"{NAMES_JSON.relative_to(ROOT)}: {len(sorted_names)} 件")
    return 0


if __name__ == "__main__":
    sys.exit(main())
